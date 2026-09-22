import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import math2docx
from . import css2json

def style_parser(css_path):
    """
    Reads a CSS file and converts it into a style dictionary using css2json.
    """
    try:
        return css2json.parse_css_file(css_path)
    except Exception as e:
        print(f"Warning: Could not parse CSS ({e})")
        return {}

def markdown_parser(md_text):
    """
    Parses Markdown text into a list of structured blocks.
    Supported blocks: heading, paragraph, list, table, math
    """
    blocks = []
    lines = md_text.split('\n')
    
    current_table = []
    
    for line in lines:
        line_s = line.strip()
        
        # Table Parsing
        if line_s.startswith('|') and line_s.endswith('|'):
            current_table.append(line_s)
            continue
        elif current_table:
            blocks.append(('table', current_table))
            current_table = []
            
        if not line_s:
            continue
            
        # Heading Parsing
        if line_s.startswith('#'):
            level = len(line_s) - len(line_s.lstrip('#'))
            text = line_s.lstrip('#').strip()
            blocks.append(('heading', level, text))
            
        # List Parsing
        elif line_s.startswith('- ') or line_s.startswith('* '):
            text = line_s[2:].strip()
            blocks.append(('list', text, 'unordered'))
            
        elif re.match(r'^\d+\.\s+', line_s):
            m = re.match(r'^\d+\.\s+(.*)', line_s)
            blocks.append(('list', m.group(1), 'ordered'))
            
        # Math Block Parsing
        elif line_s.startswith('$$') and line_s.endswith('$$'):
            blocks.append(('math', line_s.strip('$').strip()))
            
        # Standard Paragraph
        else:
            blocks.append(('paragraph', line_s))
            
    if current_table:
        blocks.append(('table', current_table))
        
    return blocks

def _apply_css_to_run(run, css_rules):
    """Helper to apply CSS rules to a docx run or font."""
    font = run.font if hasattr(run, 'font') else run
    
    if 'color' in css_rules:
        hex_col = css_rules['color'].replace('#', '').strip()
        if len(hex_col) == 6:
            try:
                r, g, b = tuple(int(hex_col[i:i+2], 16) for i in (0, 2, 4))
                font.color.rgb = RGBColor(r, g, b)
            except ValueError:
                pass
                
    if 'font-size' in css_rules:
        size = css_rules['font-size'].replace('pt', '').strip()
        try:
            font.size = Pt(float(size))
        except ValueError:
            pass
            
    if 'font-family' in css_rules:
        font.name = css_rules['font-family'].strip("'\"")
        
    if 'font-weight' in css_rules:
        font.bold = (css_rules['font-weight'].lower() == 'bold')
        
    if 'font-style' in css_rules:
        font.italic = (css_rules['font-style'].lower() == 'italic')
        
    if 'text-decoration' in css_rules:
        dec = css_rules['text-decoration'].lower()
        if dec == 'underline':
            font.underline = True
        elif dec == 'line-through':
            font.strike = True
        elif dec == 'none':
            font.underline = False
            font.strike = False
            
    if 'text-transform' in css_rules:
        trans = css_rules['text-transform'].lower()
        if trans == 'uppercase':
            font.all_caps = True
        elif trans == 'small-caps':
            font.small_caps = True
            
    if 'background-color' in css_rules:
        # Maps text background to highlight color index
        from docx.enum.text import WD_COLOR_INDEX
        color_map = {
            'auto': WD_COLOR_INDEX.AUTO, 'black': WD_COLOR_INDEX.BLACK, 'blue': WD_COLOR_INDEX.BLUE,
            'bright-green': WD_COLOR_INDEX.BRIGHT_GREEN, 'dark-blue': WD_COLOR_INDEX.DARK_BLUE,
            'dark-red': WD_COLOR_INDEX.DARK_RED, 'dark-yellow': WD_COLOR_INDEX.DARK_YELLOW,
            'gray-25': WD_COLOR_INDEX.GRAY_25, 'gray-50': WD_COLOR_INDEX.GRAY_50,
            'green': WD_COLOR_INDEX.GREEN, 'pink': WD_COLOR_INDEX.PINK, 'red': WD_COLOR_INDEX.RED,
            'teal': WD_COLOR_INDEX.TEAL, 'turquoise': WD_COLOR_INDEX.TURQUOISE,
            'violet': WD_COLOR_INDEX.VIOLET, 'white': WD_COLOR_INDEX.WHITE, 'yellow': WD_COLOR_INDEX.YELLOW
        }
        bg_col = css_rules['background-color'].lower()
        if bg_col in color_map:
            font.highlight_color = color_map[bg_col]

def _apply_paragraph_formatting(p, css_rules):
    """Helper to apply CSS rules to paragraph formatting."""
    fmt = p.paragraph_format
    
    if 'text-align' in css_rules:
        align = css_rules['text-align'].lower()
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'right':
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == 'justify':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == 'distribute':
            p.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
        elif align == 'left':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
    if 'margin-top' in css_rules:
        val = css_rules['margin-top'].replace('pt', '').strip()
        try:
            fmt.space_before = Pt(float(val))
        except ValueError:
            pass
            
    if 'margin-bottom' in css_rules:
        val = css_rules['margin-bottom'].replace('pt', '').strip()
        try:
            fmt.space_after = Pt(float(val))
        except ValueError:
            pass
            
    if 'line-height' in css_rules:
        try:
            fmt.line_spacing = float(css_rules['line-height'])
        except ValueError:
            pass
            
    if 'page-break-before' in css_rules:
        if css_rules['page-break-before'].lower() == 'always':
            fmt.page_break_before = True

import docx.opc.constants
from docx.oxml.shared import OxmlElement, qn
import urllib.request
import io
import random

def _add_internal_hyperlink(paragraph, text, anchor):
    """Adds a clickable internal hyperlink to a bookmark anchor in the docx."""
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), anchor)
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    # Make it superscript since it's usually a footnote marker
    vertAlign = OxmlElement('w:vertAlign')
    vertAlign.set(qn('w:val'), 'superscript')
    rPr.append(vertAlign)
    
    # Make it blue
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0563C1')
    rPr.append(c)
    
    new_run.append(rPr)
    
    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)
    
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def _add_bookmark(paragraph, text, anchor):
    """Creates a bookmark anchor in the document and places text inside it."""
    # We need a random ID to prevent collisions
    bm_id = str(random.randint(10000, 99999))
    
    bm_start = OxmlElement('w:bookmarkStart')
    bm_start.set(qn('w:id'), bm_id)
    bm_start.set(qn('w:name'), anchor)
    paragraph._p.append(bm_start)
    
    run = paragraph.add_run(text)
    
    bm_end = OxmlElement('w:bookmarkEnd')
    bm_end.set(qn('w:id'), bm_id)
    paragraph._p.append(bm_end)
    return run

def _add_hyperlink(paragraph, text, url, styles, tag):
    """Adds a real hyperlink to a docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    c = OxmlElement('w:color')
    if 'hyperlink' in styles and 'color' in styles['hyperlink']:
        color_hex = styles['hyperlink']['color'].replace('#', '').strip()
        c.set(qn('w:val'), color_hex)
    else:
        c.set(qn('w:val'), '0563C1')
    rPr.append(c)
    
    u = OxmlElement('w:u')
    if 'hyperlink' in styles and 'text-decoration' in styles['hyperlink'] and styles['hyperlink']['text-decoration'] == 'none':
        pass
    else:
        u.set(qn('w:val'), 'single')
        rPr.append(u)
        
    new_run.append(rPr)
    
    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)
    
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def _add_image(paragraph, url, caption, styles):
    """Adds an inline image fetched from URL."""
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            image_stream = io.BytesIO(response.read())
        
        run = paragraph.add_run()
        
        # Apply CSS dimensions if available
        width = None
        if 'image' in styles and 'width' in styles['image']:
            w_str = styles['image']['width']
            if 'in' in w_str: width = docx.shared.Inches(float(w_str.replace('in', '').strip()))
            
        if width:
            run.add_picture(image_stream, width=width)
        else:
            run.add_picture(image_stream, width=docx.shared.Inches(4))
            
    except Exception as e:
        paragraph.add_run(f"[Image Failed: {url}]")

def _process_inline(paragraph, text, styles, tag):
    """Parses text for images, links, math, and bold/italic."""
    # Simple regex to split by math, image, or link. Captures the match as a group.
    # Group 5: Footnote definition [^1]: ...
    # Group 6: Footnote annotation [^1]
    pattern = r'(\$\$.*?\$\$)|(!\[.*?\]\(.*?\))|(\[.*?\]\(.*?\))|(<span.*?>.*?</span>)|(\[\^.*?\]:.*?$)|(\[\^.*?\])'
    parts = re.split(pattern, text)
    
    for part in parts:
        if not part: continue
        
        if part.startswith('$$') and part.endswith('$$'):
            math_text = part.strip('$').strip()
            try:
                math2docx.add_math(paragraph, math_text)
            except Exception:
                run = paragraph.add_run(part)
                if tag in styles: _apply_css_to_run(run, styles[tag])
                
        elif part.startswith('![') and part.endswith(')'):
            m = re.match(r'!\[(.*?)\]\((.*?)\)', part)
            if m:
                _add_image(paragraph, m.group(2), m.group(1), styles)
                
        elif part.startswith('[') and part.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if m:
                _add_hyperlink(paragraph, m.group(1), m.group(2), styles, tag)
                
        elif part.startswith('<span') and part.endswith('</span>'):
            m = re.match(r'<span.*?>(.*?)</span>', part)
            if m:
                caption_text = m.group(1)
                run = paragraph.add_run(caption_text)
                run.italic = True
                paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                if tag in styles: _apply_css_to_run(run, styles[tag])
                
        elif part.startswith('[^') and ']:' in part:
            # Footnote definition e.g. [^1]: The text
            m = re.match(r'\[\^(.*?)\]:\s*(.*)', part)
            if m:
                fn_id = m.group(1)
                fn_text = m.group(2)
                # Create a bookmark here for the annotation to jump to
                run = _add_bookmark(paragraph, f"[{fn_id}]: {fn_text}", f"footnote_{fn_id}")
                if tag in styles: _apply_css_to_run(run, styles[tag])
                
        elif part.startswith('[^') and part.endswith(']'):
            # Footnote annotation e.g. [^1]
            m = re.match(r'\[\^(.*?)\]', part)
            if m:
                fn_id = m.group(1)
                _add_internal_hyperlink(paragraph, f"[{fn_id}]", f"footnote_{fn_id}")
                
        else:
            # Handle plain text + bold/italic
            subparts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', part)
            for subpart in subparts:
                if not subpart: continue
                if subpart.startswith('**') and subpart.endswith('**'):
                    run = paragraph.add_run(subpart[2:-2])
                    run.bold = True
                    if tag in styles: _apply_css_to_run(run, styles[tag])
                elif subpart.startswith('*') and subpart.endswith('*'):
                    run = paragraph.add_run(subpart[1:-1])
                    run.italic = True
                    if tag in styles: _apply_css_to_run(run, styles[tag])
                else:
                    run = paragraph.add_run(subpart)
                    if tag in styles: _apply_css_to_run(run, styles[tag])

def docx_writer(blocks: list, styles: dict, output_file: str):
    """
    Translates parsed markdown blocks into a Word document and applies CSS styling.
    """
    doc = Document()
    
    # Global body style
    if 'body' in styles:
        normal_style = doc.styles['Normal']
        _apply_css_to_run(normal_style, styles['body'])
        _apply_paragraph_formatting(normal_style, styles['body'])

    for block in blocks:
        btype = block[0]
        
        if btype == 'heading':
            level, text = block[1], block[2]
            p = doc.add_heading('', level=level)
            _process_inline(p, text, styles, f'h{level}')
            if f'h{level}' in styles:
                _apply_paragraph_formatting(p, styles[f'h{level}'])
                
        elif btype == 'paragraph':
            text = block[1]
            p = doc.add_paragraph()
            if 'p' in styles:
                _apply_paragraph_formatting(p, styles['p'])
            _process_inline(p, text, styles, 'p')
                
        elif btype == 'list':
            text = block[1]
            list_type = block[2] if len(block) > 2 else 'unordered'
            style = 'List Number' if list_type == 'ordered' else 'List Bullet'
            p = doc.add_paragraph(style=style)
            if 'list' in styles:
                _apply_paragraph_formatting(p, styles['list'])
            _process_inline(p, text, styles, 'list')
                
        elif btype == 'math':
            math_text = block[1]
            p = doc.add_paragraph()
            try:
                math2docx.add_math(p, math_text)
                if 'equation' in styles:
                    _apply_paragraph_formatting(p, styles['equation'])
            except Exception:
                p.add_run(f"$$ {math_text} $$")
                
        elif btype == 'table':
            table_lines = block[1]
            rows = [
                [cell.strip() for cell in line.strip('|').split('|')]
                for line in table_lines if '---' not in line
            ]
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for r_idx, row in enumerate(rows):
                    for c_idx, val in enumerate(row):
                        table.cell(r_idx, c_idx).text = val
                if 'table' in styles:
                    pass

    doc.save(output_file)


if __name__ == "__main__":
    # Example Usage:
    styles = style_parser('engine/style.css')
    md_blocks = markdown_parser('## Hello World\nThis is a **bold** test.')
    docx_writer(md_blocks, styles, 'output.docx')
    