import math2docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

# Initialize the document
doc = Document()

# ---------------------------------------------------------
# 1. GLOBAL STYLE MODIFICATION
# ---------------------------------------------------------
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# ---------------------------------------------------------
# 2. HEADINGS & BASIC PARAGRAPHS
# ---------------------------------------------------------
doc.add_heading('Comprehensive python-docx Test Script', level=1)
doc.add_paragraph(
    'This document tests various features including basic structure, '
    'direct styling, tables with shading, and LaTeX math equations.'
)

# ---------------------------------------------------------
# 3. DIRECT FORMATTING & ALIGNMENT
# ---------------------------------------------------------
p_aligned = doc.add_paragraph()
p_aligned.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_center = p_aligned.add_run('This paragraph is centered and custom-styled.')
run_center.font.name = 'Arial'
run_center.font.size = Pt(13)
run_center.font.bold = True
run_center.font.color.rgb = RGBColor(0, 102, 204)  # Blue color

# Mixed inline formatting (bold/italic runs)
p_mixed = doc.add_paragraph('Standard text with ')
p_mixed.add_run('bold text').bold = True
p_mixed.add_run(' and ')
p_mixed.add_run('italic text').italic = True
p_mixed.add_run('.')

# Lists and built-in styles
doc.add_heading('Lists and Built-in Styles', level=2)
doc.add_paragraph('First item in a bullet list', style='List Bullet')
doc.add_paragraph('Second item in a bullet list', style='List Bullet')
doc.add_paragraph(
    'This uses the built-in Intense Quote style.', style='Intense Quote'
)

# ---------------------------------------------------------
# 4. TABLES & CELL SHADING
# ---------------------------------------------------------
doc.add_heading('Tables with Styling and Shading', level=2)
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'

# Populate table cells
table.cell(0, 0).text = 'Header Cell 1'
table.cell(0, 1).text = 'Header Cell 2'
table.cell(1, 0).text = 'Data Row 1'
table.cell(1, 1).text = 'Data Row 2'

# Apply background shading (light gray) to the first header cell via XML
cell = table.cell(0, 0)
shading_elm = parse_xml(r'<w:shd {} w:fill="D3D3D3"/>'.format(nsdecls('w')))
cell._tc.get_or_add_tcPr().append(shading_elm)

# ---------------------------------------------------------
# 5. LATEX MATH EQUATIONS
# ---------------------------------------------------------
doc.add_heading('LaTeX Math Equations', level=2)
p_math = doc.add_paragraph('Here is a native Word equation generated from LaTeX: ')
math2docx.add_math(p_math, r'x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}')

# ---------------------------------------------------------
# 6. SAVE DOCUMENT
# ---------------------------------------------------------
doc.save('comprehensive_test.docx')
print(
    'Successfully generated and saved test document as'
    ' "comprehensive_test.docx"!'
)